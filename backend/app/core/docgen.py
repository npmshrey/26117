"""Generates real DOCX deliverables from agent output — actual files on
disk, not placeholders."""
import time
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor

ARTIFACTS_DIR = Path(__file__).resolve().parent.parent.parent / "storage" / "artifacts"
ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)


def build_approval_note(task_id: str, prompt: str, findings: str, citations: list[dict], model_id: str) -> Path:
    doc = Document()

    title = doc.add_heading("Approval Note", level=1)
    title.runs[0].font.color.rgb = RGBColor(0x1A, 0x12, 0x06)

    meta = doc.add_paragraph()
    meta.add_run(f"Task ID: {task_id}\n").italic = True
    meta.add_run(f"Generated: {time.strftime('%Y-%m-%d %H:%M:%S')}\n").italic = True
    meta.add_run(f"Model: {model_id}\n").italic = True

    doc.add_heading("Request", level=2)
    doc.add_paragraph(prompt)

    doc.add_heading("Findings", level=2)
    for line in findings.split("\n"):
        line = line.strip()
        if line:
            doc.add_paragraph(line, style="List Bullet" if not line[0].isdigit() else None)

    if citations:
        doc.add_heading("Evidence & Citations", level=2)
        for c in citations:
            p = doc.add_paragraph()
            p.add_run(f"{c.get('doc_name', 'unknown')} (page {c.get('page', '?')}): ").bold = True
            p.add_run(c.get("snippet", ""))

    doc.add_heading("Sovereignty Note", level=2)
    doc.add_paragraph(
        "This document was generated entirely on-premise. No content was "
        "transmitted to an external network during generation."
    ).italic = True

    out_path = ARTIFACTS_DIR / f"approval_note_{task_id}.docx"
    doc.save(out_path)
    return out_path
